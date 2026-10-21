import io,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image
from formal_pdf import brand_logo,LABELS

def safe(value):
 if value is None: return ''
 if isinstance(value,list): value='; '.join(' | '.join(map(str,row)) if isinstance(row,list) else str(row) for row in value)
 if isinstance(value,dict): value='; '.join(f'{k}: {v}' for k,v in value.items())
 if isinstance(value,str) and value.startswith(('=','+','-','@')): return "'"+value
 return value

def export_excel(rows,company,branding,root,users,title):
 book=Workbook(); main=book.active; main.title='Reporte'; details=book.create_sheet('Detalle de campos')
 name=company.get('nombre_comercial') or company.get('razon_social') or 'QualityLite'
 logo=brand_logo(root,branding)
 def prepare(sheet,headers,widths):
  sheet.sheet_view.showGridLines=False
  last=max(len(headers),4); sheet.merge_cells(start_row=1,start_column=2,end_row=1,end_column=last); sheet.cell(1,2,safe(name)).font=Font(name='Calibri',size=16,bold=True,color='173454')
  sheet.merge_cells(start_row=2,start_column=2,end_row=2,end_column=last); sheet.cell(2,2,safe(title)).font=Font(name='Calibri',size=12,color='65788C')
  sheet.row_dimensions[1].height=27; sheet.row_dimensions[2].height=22
  sheet.cell(3,2,safe(branding.get('department','Sistema de Gestión de Calidad e Inocuidad'))).font=Font(size=9,color='65788C')
  if logo:
   picture=Image(str(logo)); ratio=min(70/picture.width,44/picture.height); picture.width*=ratio; picture.height*=ratio; sheet.add_image(picture,'A1')
  for index,header in enumerate(headers,1):
   cell=sheet.cell(5,index,header); cell.fill=PatternFill('solid',fgColor='173454'); cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(wrap_text=True,vertical='center')
   sheet.column_dimensions[cell.column_letter].width=widths[index-1]
  sheet.row_dimensions[5].height=28; sheet.freeze_panes='A6'; sheet.print_title_rows='1:5'; sheet.sheet_properties.pageSetUpPr.fitToPage=True
  sheet.page_setup.orientation='landscape'; sheet.page_setup.paperSize=sheet.PAPERSIZE_LETTER; sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0
  sheet.oddFooter.center.text='Página &P de &N'; sheet.oddFooter.left.text='QualityLite | Documento de consulta'
 prepare(main,['Código / folio','Título','Versión','Estado','Área','Proceso','Responsable','Fecha'],[24,45,10,24,25,28,28,22])
 prepare(details,['Folio','Sección','Campo','Respuesta'],[26,32,38,80])
 for row in rows:
  data=json.loads(row['data']); main.append([safe(v) for v in (row['code'],row['title'],row['version'],row['state'],data.get('area',''),data.get('proceso',''),users.get(row['owner'],''),row['created'][:19])])
  if data.get('schema'):
   for field in data['schema']: details.append([safe(v) for v in (row['code'],field.get('section','Datos del formato'),field['label'],data.get('answers',{}).get(field['id'],''))])
  else:
   for key,value in data.items():
    if key.startswith('_') or key in ('state_dates','closed_by','closed_at'): continue
    details.append([safe(row['code']),'Información del documento',safe(LABELS.get(key,key.replace('_',' ').capitalize())),safe(value)])
 for sheet in book:
  sheet.auto_filter.ref=f'A5:{sheet.cell(max(sheet.max_row,5),sheet.max_column).coordinate}'
  for cells in sheet.iter_rows(min_row=6):
   for cell in cells:
    cell.alignment=Alignment(vertical='top',wrap_text=True); cell.font=Font(name='Calibri',size=11,color='23364A'); cell.border=Border(bottom=Side(style='hair',color='DCE4EC'))
   sheet.row_dimensions[cells[0].row].height=42
 out=io.BytesIO(); book.save(out); return out.getvalue()
